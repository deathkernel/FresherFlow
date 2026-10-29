import io
import re
import secrets
import sqlite3
from datetime import date, datetime, timedelta
from zipfile import BadZipFile, LargeZipFile, ZipFile

from flask import (
    Blueprint,
    current_app,
    flash,
    redirect,
    render_template,
    request,
    send_from_directory,
    session,
    url_for,
)
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from database.database import get_db
from database.vacancy_deduplication import find_duplicate_vacancy, new_vacancy_code, vacancy_fingerprint
from routes.decorators import role_required
from security import valid_website_url

employer_bp = Blueprint("employer", __name__, url_prefix="/employer")
VACANCY_TYPES = {"Internship", "Entry-level Job"}
VACANCY_STATUSES = {"draft", "active", "closed"}
APPLICATION_STATUSES = {"Applied", "Shortlisted", "Selected", "Rejected"}
EXCEL_HEADERS = {
    "title": "title", "type": "vacancy_type", "vacancy type": "vacancy_type", "vacancy_type": "vacancy_type",
    "location": "location", "salary": "salary", "salary / stipend": "salary", "salary/stipend": "salary",
    "stipend": "salary", "deadline": "deadline", "skills": "skills", "eligibility": "eligibility",
    "description": "description", "role description": "description",
}
MAX_BULK_VACANCIES = 100
MAX_EXCEL_UNCOMPRESSED_BYTES = 25 * 1024 * 1024
MAX_EXCEL_MEMBER_BYTES = 10 * 1024 * 1024
MAX_EXCEL_ARCHIVE_ENTRIES = 1000


def experience_requirement_invalid(vacancy_type, eligibility):
    if vacancy_type != "Entry-level Job":
        return False
    text = (eligibility or "").strip().lower()
    if not text:
        return False

    negative_experience = re.search(
        r"\b(?:no|without|zero|0)\b.{0,20}\bexperience\b", text
    )
    if negative_experience:
        return False

    terms = (
        "work experience",
        "professional experience",
        "prior experience",
        "previous experience",
        "relevant experience",
        "industry experience",
        "experience required",
        "experience mandatory",
        "minimum experience",
    )
    if any(term in text for term in terms):
        return True

    experience_pattern = re.compile(
        r"\b(?:minimum\s+)?\d+(?:\.\d+)?\s*\+?\s*(?:year|years|yr|yrs)\b(?:\s+of)?\s+experience\b"
    )
    reverse_pattern = re.compile(
        r"\bexperience\b.{0,30}\b\d+(?:\.\d+)?\s*\+?\s*(?:year|years|yr|yrs)\b"
    )
    return bool(experience_pattern.search(text) or reverse_pattern.search(text))


def deadline_invalid(deadline):
    if not deadline:
        return False

    try:
        return date.fromisoformat(deadline) < date.today()
    except ValueError:
        return True


def vacancy_form(form):
    title = form.get("title", "").strip()
    vacancy_type = form.get("vacancy_type", "").strip()
    location = form.get("location", "").strip()
    description = form.get("description", "").strip()
    eligibility = form.get("eligibility", "").strip()
    deadline = form.get("deadline") or None

    if not title or not location or not description:
        return None, "Title, location and description are required."
    if vacancy_type not in VACANCY_TYPES:
        return None, "Choose a valid vacancy type."
    if experience_requirement_invalid(vacancy_type, eligibility):
        return (
            None,
            "Entry-level Jobs cannot require prior work experience. "
            "For roles requiring experience, publish an Internship instead.",
        )
    if deadline_invalid(deadline):
        return None, "Application deadline must be today or a future date."

    return {
        "title": title,
        "vacancy_type": vacancy_type,
        "description": description,
        "location": location,
        "salary": form.get("salary", "").strip(),
        "skills": form.get("skills", "").strip(),
        "eligibility": eligibility,
        "deadline": deadline,
    }, None


def normalize_excel_header(value):
    return " ".join(
        str(value or "").strip().lower().replace("_", " ").split()
    )


def excel_cell_text(value):
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def bulk_vacancy_forms(form):
    titles = form.getlist("title")
    types = form.getlist("vacancy_type")
    locations = form.getlist("location")
    descriptions = form.getlist("description")
    salaries = form.getlist("salary")
    skills = form.getlist("skills")
    eligibilities = form.getlist("eligibility")
    deadlines = form.getlist("deadline")

    count = len(titles)
    if count > MAX_BULK_VACANCIES:
        return None, f"Import at most {MAX_BULK_VACANCIES} vacancies per batch."
    if (
        not count
        or count != len(types)
        or count != len(locations)
        or count != len(descriptions)
    ):
        return None, "Please complete each vacancy card before publishing."

    vacancies = []
    for index in range(count):
        data = {
            "title": titles[index].strip(),
            "vacancy_type": types[index].strip(),
            "location": locations[index].strip(),
            "description": descriptions[index].strip(),
            "salary": salaries[index].strip() if index < len(salaries) else "",
            "skills": skills[index].strip() if index < len(skills) else "",
            "eligibility": (
                eligibilities[index].strip()
                if index < len(eligibilities)
                else ""
            ),
            "deadline": (deadlines[index].strip() or None) if index < len(deadlines) else None,
        }

        number = index + 1
        if not data["title"] or not data["location"] or not data["description"]:
            return (
                None,
                f"Vacancy {number}: title, location and description are required.",
            )
        if data["vacancy_type"] not in VACANCY_TYPES:
            return None, f"Vacancy {number}: choose a valid vacancy type."
        if experience_requirement_invalid(
            data["vacancy_type"], data["eligibility"]
        ):
            return (
                None,
                f"Vacancy {number}: Entry-level Jobs cannot require prior work experience.",
            )
        if deadline_invalid(data["deadline"]):
            return (
                None,
                f"Vacancy {number}: application deadline must be today "
                "or a future date.",
            )

        vacancies.append(data)

    return vacancies, None


def excel_vacancy_forms(file_storage):
    filename = (file_storage.filename or "").strip().lower()
    if not filename.endswith(".xlsx"):
        return None, "Please upload an Excel .xlsx file."

    raw_file = file_storage.read()
    try:
        # XLSX is a ZIP container. Check expanded sizes before XML parsing so a
        # small compressed upload cannot expand into an unbounded workbook.
        with ZipFile(io.BytesIO(raw_file)) as archive:
            infos = archive.infolist()
            if (
                len(infos) > MAX_EXCEL_ARCHIVE_ENTRIES
                or any(info.file_size > MAX_EXCEL_MEMBER_BYTES for info in infos)
                or sum(info.file_size for info in infos) > MAX_EXCEL_UNCOMPRESSED_BYTES
            ):
                return (
                    None,
                    "The Excel file expands beyond the allowed safety limits.",
                )
        workbook = load_workbook(
            filename=io.BytesIO(raw_file),
            read_only=True,
            data_only=True,
        )
    except (InvalidFileException, BadZipFile, LargeZipFile, OSError, ValueError, EOFError):
        return (
            None,
            "The Excel file could not be read. Please upload a valid .xlsx file.",
        )

    try:
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)

        try:
            header_row = next(rows)
        except StopIteration:
            return None, "The Excel file is empty."

        headers = {}
        for index, value in enumerate(header_row):
            normalized = normalize_excel_header(value)
            if normalized in EXCEL_HEADERS:
                headers[EXCEL_HEADERS[normalized]] = index

        required = {"title", "vacancy_type", "location", "description"}
        missing = required - set(headers)
        if missing:
            return (
                None,
                "Missing required Excel columns: "
                + ", ".join(sorted(missing))
                + ".",
            )

        vacancies = []
        for row_number, row in enumerate(rows, start=2):
            if not any(value not in (None, "") for value in row):
                continue

            data = {}
            for field in EXCEL_HEADERS.values():
                index = headers.get(field)
                if index is not None and index < len(row):
                    data[field] = excel_cell_text(row[index])
                else:
                    data[field] = ""

            data["deadline"] = data["deadline"] or None

            if not data["title"] or not data["location"] or not data["description"]:
                return (
                    None,
                    f"Excel row {row_number}: title, location and description "
                    "are required.",
                )
            if data["vacancy_type"] not in VACANCY_TYPES:
                return (
                    None,
                    f"Excel row {row_number}: type must be Internship "
                    "or Entry-level Job.",
                )
            if experience_requirement_invalid(
                data["vacancy_type"], data["eligibility"]
            ):
                return (
                    None,
                    f"Excel row {row_number}: Entry-level Jobs cannot require "
                    "prior work experience.",
                )
            if deadline_invalid(data["deadline"]):
                return (
                    None,
                    f"Excel row {row_number}: application deadline must be "
                    "today or a future date.",
                )

            if len(vacancies) >= MAX_BULK_VACANCIES:
                return (
                    None,
                    f"Import at most {MAX_BULK_VACANCIES} vacancies per batch.",
                )
            vacancies.append(data)

        if not vacancies:
            return None, "No vacancy rows were found in the Excel file."

        return vacancies, None
    finally:
        workbook.close()


@employer_bp.get("/dashboard")
@role_required("employer")
def dashboard():
    db = get_db()
    uid = session["user_id"]

    stats = {
        "live_jobs": db.execute(
            "SELECT COUNT(*) c FROM vacancies v "
            "JOIN employer_profiles ep ON ep.user_id=v.employer_id "
            "WHERE v.employer_id=? AND v.status='active' "
            "AND v.moderation_status='approved' AND ep.account_status='active' "
            "AND ep.verification_status='verified' "
            "AND (v.deadline IS NULL OR trim(v.deadline)='' OR v.deadline>=date('now','localtime'))",
            (uid,),
        ).fetchone()["c"],
        "applications": db.execute(
            "SELECT COUNT(*) c FROM applications a "
            "JOIN vacancies v ON v.id=a.vacancy_id "
            "WHERE v.employer_id=?",
            (uid,),
        ).fetchone()["c"],
        "shortlisted": db.execute(
            "SELECT COUNT(*) c FROM applications a "
            "JOIN vacancies v ON v.id=a.vacancy_id "
            "WHERE v.employer_id=? AND a.status='Shortlisted'",
            (uid,),
        ).fetchone()["c"],
        "selected": db.execute(
            "SELECT COUNT(*) c FROM applications a "
            "JOIN vacancies v ON v.id=a.vacancy_id "
            "WHERE v.employer_id=? AND a.status='Selected'",
            (uid,),
        ).fetchone()["c"],
    }

    posted_vacancies = db.execute(
        "SELECT id,title,vacancy_type,location,status,moderation_status,deadline,vacancy_code "
        "FROM vacancies WHERE employer_id=? ORDER BY id DESC LIMIT 8",
        (uid,),
    ).fetchall()
    applications = db.execute(
        "SELECT a.*,v.title,u.name,u.email "
        "FROM applications a "
        "JOIN vacancies v ON v.id=a.vacancy_id "
        "JOIN users u ON u.id=a.student_id "
        "WHERE v.employer_id=? ORDER BY a.id DESC LIMIT 8",
        (uid,),
    ).fetchall()

    return render_template(
        "employer/dashboard.html",
        stats=stats,
        posted_vacancies=posted_vacancies,
        applications=applications,
    )


@employer_bp.route("/profile", methods=["GET", "POST"])
@role_required("employer")
def profile():
    db = get_db()
    uid = session["user_id"]

    if request.method == "POST":
        data = [
            request.form.get(key, "").strip()
            for key in (
                "organization_name",
                "organization_type",
                "website",
                "location",
                "description",
            )
        ]
        if not data[0]:
            flash("Organization name is required.","error")
            return redirect(url_for("employer.profile"))
        if not valid_website_url(data[2]):
            flash("Website must be a valid http:// or https:// URL.","error")
            return redirect(url_for("employer.profile"))
        existing = db.execute(
            "SELECT organization_name, organization_type, website, location, description, verification_status "
            "FROM employer_profiles WHERE user_id=?",
            (uid,),
        ).fetchone()
        if not existing:
            return render_template("404.html"), 404

        identity_changed = any(
            data[index] != (existing[field] or "")
            for index, field in enumerate(
                ("organization_name", "organization_type", "website", "location", "description")
            )
        )
        if identity_changed:
            db.execute(
                "UPDATE employer_profiles SET organization_name=?, organization_type=?, website=?, location=?, "
                "description=?, verification_status='pending', verification_note=?, verified_at=NULL WHERE user_id=?",
                (*data, "Employer details changed; re-verification required.", uid),
            )
            flash("Profile updated. Employer verification is pending review again.", "info")
        else:
            db.execute(
                "UPDATE employer_profiles SET organization_name=?, organization_type=?, website=?, location=?, description=? WHERE user_id=?",
                (*data, uid),
            )
        db.commit()
        flash("Organization profile updated.", "success")
        return redirect(url_for("employer.profile"))

    profile = db.execute(
        "SELECT * FROM employer_profiles WHERE user_id=?",
        (uid,),
    ).fetchone()
    stats = {
        "jobs": db.execute(
            "SELECT COUNT(*) FROM vacancies WHERE employer_id=?", (uid,)
        ).fetchone()[0],
        "live_jobs": db.execute(
            "SELECT COUNT(*) FROM vacancies v "
            "JOIN employer_profiles ep ON ep.user_id=v.employer_id "
            "WHERE v.employer_id=? AND v.status='active' AND v.moderation_status='approved' "
            "AND ep.account_status='active' AND ep.verification_status='verified' "
            "AND (v.deadline IS NULL OR trim(v.deadline)='' OR v.deadline>=date('now','localtime'))",
            (uid,),
        ).fetchone()[0],
        "pending_jobs": db.execute(
            "SELECT COUNT(*) FROM vacancies WHERE employer_id=? AND moderation_status='pending'",
            (uid,),
        ).fetchone()[0],
        "draft_jobs": db.execute(
            "SELECT COUNT(*) FROM vacancies WHERE employer_id=? AND status='draft'",
            (uid,),
        ).fetchone()[0],
        "applications": db.execute(
            "SELECT COUNT(*) FROM applications a JOIN vacancies v ON v.id=a.vacancy_id WHERE v.employer_id=?",
            (uid,),
        ).fetchone()[0],
    }
    return render_template("employer/profile.html", profile=profile, stats=stats)


@employer_bp.route("/vacancies/new", methods=["GET", "POST"])
@role_required("employer")
def new_vacancy():
    if request.method == "POST":
        data, error = vacancy_form(request.form)
        if error:
            flash(error, "error")
            return render_template("employer/post-vacancy.html")

        db = get_db()
        uid = session["user_id"]
        duplicate = find_duplicate_vacancy(db, uid, data)
        if duplicate:
            flash(f"This looks like a duplicate of {duplicate['vacancy_code']} ({duplicate['title']}). No second vacancy was created.", "warning")
            return redirect(url_for("employer.vacancies"))
        publish = request.form.get("publish") == "1"
        try:
            status = "active" if publish else "draft"
            moderation_status = "pending"
            db.execute(
                "INSERT INTO vacancies(employer_id,title,vacancy_type,description,location,salary,skills,eligibility,deadline,status,moderation_status,vacancy_code,dedupe_fingerprint) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)",
                (uid, data["title"], data["vacancy_type"], data["description"], data["location"], data["salary"], data["skills"], data["eligibility"], data["deadline"], status, moderation_status, new_vacancy_code(), vacancy_fingerprint(uid, data)),
            )
            db.commit()
        except sqlite3.IntegrityError:
            db.rollback()
            flash("The vacancy could not be saved.", "error")
            return render_template("employer/post-vacancy.html")
        flash("Vacancy submitted for admin approval." if publish else "Vacancy saved as draft.","success")
        return redirect(url_for("employer.vacancies"))
    return render_template("employer/post-vacancy.html")


@employer_bp.route("/vacancies/bulk-new", methods=["GET", "POST"])
@role_required("employer")
def bulk_new_vacancies():
    if request.method == "GET":
        return render_template("employer/bulk-vacancies.html")

    if request.form.get("import_excel") == "1":
        excel_file = request.files.get("excel_file")
        if not excel_file or not excel_file.filename:
            flash("Choose an Excel .xlsx file first.", "error")
            return render_template("employer/bulk-vacancies.html"), 400

        vacancies, error = excel_vacancy_forms(excel_file)
    else:
        vacancies, error = bulk_vacancy_forms(request.form)

    if error:
        flash(error, "error")
        return render_template("employer/bulk-vacancies.html"), 400

    db = get_db()
    uid = session["user_id"]

    inserted = 0
    skipped = []
    batch_fingerprints = set()
    try:
        for data in vacancies:
            fingerprint = vacancy_fingerprint(uid, data)
            duplicate = find_duplicate_vacancy(db, uid, data)
            if duplicate or fingerprint in batch_fingerprints:
                code = duplicate["vacancy_code"] if duplicate else "another row in this upload"
                skipped.append(f"{data['title']} ({code})")
                continue
            db.execute(
                "INSERT INTO vacancies(employer_id,title,vacancy_type,description,location,salary,skills,eligibility,deadline,status,moderation_status,vacancy_code,dedupe_fingerprint) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)",
                (uid, data["title"], data["vacancy_type"], data["description"], data["location"], data["salary"], data["skills"], data["eligibility"], data["deadline"], "active", "pending", new_vacancy_code(), fingerprint),
            )
            inserted += 1
            batch_fingerprints.add(fingerprint)
        db.commit()
    except sqlite3.IntegrityError:
        db.rollback()
        flash(
            "None of the vacancies were imported because one or more "
            "entries were invalid.",
            "error",
        )
        return render_template("employer/bulk-vacancies.html"), 400

    message = f"{inserted} new vacancies submitted for admin approval."
    if skipped:
        message += f" Skipped {len(skipped)} duplicate(s): " + "; ".join(skipped[:5])
        if len(skipped) > 5:
            message += f"; and {len(skipped) - 5} more"
    flash(message, "success" if not skipped else "warning")
    return redirect(url_for("employer.vacancies"))


@employer_bp.get("/vacancies")
@role_required("employer")
def vacancies():
    db = get_db()
    page = max(request.args.get("page", 1, type=int), 1)
    per_page = 20
    offset = (page - 1) * per_page
    uid = session["user_id"]
    total = db.execute(
        "SELECT COUNT(*) FROM vacancies WHERE employer_id=?", (uid,)
    ).fetchone()[0]
    rows = db.execute(
        "SELECT * FROM vacancies WHERE employer_id=? ORDER BY id DESC LIMIT ? OFFSET ?",
        (uid, per_page, offset),
    ).fetchall()
    draft_count = db.execute(
        "SELECT COUNT(*) FROM vacancies WHERE employer_id=? AND status='draft'", (uid,)
    ).fetchone()[0]
    return render_template(
        "employer/vacancies.html",
        vacancies=rows,
        draft_count=draft_count,
        page=page,
        per_page=per_page,
        total=total,
    )


@employer_bp.route("/vacancies/<int:vacancy_id>/edit", methods=["GET", "POST"])
@role_required("employer")
def edit_vacancy(vacancy_id):
    db = get_db()
    job = db.execute(
        "SELECT * FROM vacancies WHERE id=? AND employer_id=?",
        (vacancy_id, session["user_id"]),
    ).fetchone()

    if not job:
        return render_template("404.html"), 404

    if request.method == "POST":
        data, error = vacancy_form(request.form)
        if error:
            flash(error, "error")
            return render_template("employer/edit-vacancy.html", job=job)
        duplicate = find_duplicate_vacancy(db, session["user_id"], data, exclude_id=vacancy_id)
        if duplicate:
            flash(f"This edit matches existing vacancy {duplicate['vacancy_code']}. Edit that vacancy instead or change these details.", "warning")
            return render_template("employer/edit-vacancy.html", job=job)
        try:
            db.execute(
                "UPDATE vacancies SET title=?, vacancy_type=?, description=?, "
                "location=?, salary=?, skills=?, eligibility=?, deadline=?, dedupe_fingerprint=?, "
                "status='draft', moderation_status='pending', "
                "moderation_note=NULL, moderated_at=NULL "
                "WHERE id=? AND employer_id=?",
                (
                    data["title"],
                    data["vacancy_type"],
                    data["description"],
                    data["location"],
                    data["salary"],
                    data["skills"],
                    data["eligibility"],
                    data["deadline"],
                    vacancy_fingerprint(session["user_id"], data),
                    vacancy_id,
                    session["user_id"],
                ),
            )
            db.commit()
        except sqlite3.IntegrityError:
            db.rollback()
            flash("The vacancy could not be updated.", "error")
            return render_template("employer/edit-vacancy.html", job=job)
        flash(
            "Vacancy updated and saved as a draft. Publish it from My Vacancies when ready.",
            "success",
        )
        return redirect(url_for("employer.vacancies"))
    return render_template("employer/edit-vacancy.html",job=job)


@employer_bp.post("/vacancies/publish-drafts")
@role_required("employer")
def publish_all_drafts():
    db = get_db()
    uid = session["user_id"]
    drafts = db.execute(
        "SELECT id, deadline FROM vacancies WHERE employer_id=? AND status='draft' ORDER BY id",
        (uid,),
    ).fetchall()
    if not drafts:
        flash("There are no draft vacancies to publish.", "error")
        return redirect(url_for("employer.vacancies"))

    invalid_deadlines = [row for row in drafts if deadline_invalid(row["deadline"])]
    if invalid_deadlines:
        flash(
            f"{len(invalid_deadlines)} draft vacancy(s) have an expired or invalid deadline. Edit them before publishing all drafts.",
            "error",
        )
        return redirect(url_for("employer.vacancies"))

    db.execute(
        "UPDATE vacancies SET status='active', moderation_status='pending', moderation_note=NULL, moderated_at=NULL WHERE employer_id=? AND status='draft'",
        (uid,),
    )
    db.commit()
    flash(
        f"{len(drafts)} draft vacancy(s) submitted for publishing and admin approval.",
        "success",
    )
    return redirect(url_for("employer.vacancies"))


@employer_bp.post("/vacancies/<int:vacancy_id>/status")
@role_required("employer")
def vacancy_status(vacancy_id):
    status = request.form.get("status")
    if status not in VACANCY_STATUSES:
        return redirect(url_for("employer.vacancies"))

    db = get_db()
    job = db.execute(
        "SELECT vacancy_type, eligibility, deadline, moderation_status "
        "FROM vacancies WHERE id=? AND employer_id=?",
        (vacancy_id, session["user_id"]),
    ).fetchone()

    if not job:
        return redirect(url_for("employer.vacancies"))

    if status == "active" and deadline_invalid(job["deadline"]):
        flash(
            "The application deadline has passed or is invalid.",
            "error",
        )
        return redirect(url_for("employer.vacancies"))
    moderation_status = "pending" if status == "active" else job["moderation_status"]
    db.execute(
        "UPDATE vacancies SET status=?, moderation_status=? "
        "WHERE id=? AND employer_id=?",
        (status, moderation_status, vacancy_id, session["user_id"]),
    )
    db.commit()
    return redirect(url_for("employer.vacancies"))


@employer_bp.get("/applications")
@role_required("employer")
def applications():
    db = get_db()
    page = max(request.args.get("page", 1, type=int), 1)
    per_page = 20
    offset = (page - 1) * per_page
    uid = session["user_id"]
    total = db.execute(
        "SELECT COUNT(*) FROM applications a JOIN vacancies v ON v.id=a.vacancy_id WHERE v.employer_id=?",
        (uid,),
    ).fetchone()[0]
    rows = db.execute("""
        SELECT a.*, v.title, u.name, u.email,
               sp.phone, sp.education, sp.college, sp.graduation_year,
               sp.skills, sp.certifications, sp.preferred_job_type,
               sp.preferred_location, sp.resume_filename, sp.profile_strength
        FROM applications a
        JOIN vacancies v ON v.id=a.vacancy_id
        JOIN users u ON u.id=a.student_id
        LEFT JOIN student_profiles sp ON sp.user_id=a.student_id
        WHERE v.employer_id=?
        ORDER BY a.id DESC
        LIMIT ? OFFSET ?
    """, (uid, per_page, offset)).fetchall()

    total_pages = (total + per_page - 1) // per_page

    return render_template(
        "employer/applications.html",
        applications=rows,
        page=page,
        per_page=per_page,
        total=total,
        total_pages=total_pages,
    )


@employer_bp.route("/applications/<int:application_id>/interview", methods=["GET", "POST"])
@role_required("employer")
def schedule_interview(application_id):
    db = get_db()
    application = db.execute(
        """SELECT a.id, a.student_id, a.status AS application_status,
                  u.name AS student_name, u.email AS student_email, v.title AS vacancy_title
           FROM applications a
           JOIN vacancies v ON v.id=a.vacancy_id
           JOIN users u ON u.id=a.student_id
           WHERE a.id=? AND v.employer_id=?""",
        (application_id, session["user_id"]),
    ).fetchone()
    if not application:
        return render_template("404.html"), 404

    if request.method == "POST":
        if application["application_status"] == "Rejected":
            flash("An interview cannot be scheduled for a rejected application.", "error")
            return redirect(url_for("employer.schedule_interview", application_id=application_id))

        raw_time = request.form.get("scheduled_at", "").strip()
        mode = request.form.get("mode", "").strip()
        meeting_url = request.form.get("meeting_url", "").strip()[:500] or None
        location = request.form.get("location", "").strip()[:250] or None
        notes = request.form.get("notes", "").strip()[:800] or None

        try:
            scheduled_time = datetime.fromisoformat(raw_time)
        except ValueError:
            flash("Choose a valid interview date and time.", "error")
            return redirect(url_for("employer.schedule_interview", application_id=application_id))

        if "T" not in raw_time or scheduled_time.tzinfo is not None:
            flash("Choose a local date and time using the interview form.", "error")
            return redirect(url_for("employer.schedule_interview", application_id=application_id))
        if scheduled_time <= datetime.now():
            flash("Interview must be scheduled for a future date and time.", "error")
            return redirect(url_for("employer.schedule_interview", application_id=application_id))
        if mode not in {"Online", "In-person", "Phone"}:
            flash("Choose a valid interview mode.", "error")
            return redirect(url_for("employer.schedule_interview", application_id=application_id))
        if meeting_url and not valid_website_url(meeting_url):
            flash("Enter a valid HTTP or HTTPS meeting link.", "error")
            return redirect(url_for("employer.schedule_interview", application_id=application_id))
        if mode == "Online" and not meeting_url:
            flash("Add a meeting link for an online interview.", "error")
            return redirect(url_for("employer.schedule_interview", application_id=application_id))
        if mode in {"In-person", "Phone"} and not location:
            flash("Add the interview location or contact details.", "error")
            return redirect(url_for("employer.schedule_interview", application_id=application_id))

        scheduled_at = scheduled_time.replace(second=0, microsecond=0).isoformat(timespec="minutes")
        try:
            db.execute("BEGIN IMMEDIATE")
            existing = db.execute(
                "SELECT id FROM interviews WHERE application_id=? AND employer_id=?",
                (application_id, session["user_id"]),
            ).fetchone()
            if existing:
                db.execute(
                    """UPDATE interviews
                       SET scheduled_at=?, mode=?, meeting_url=?, location=?, notes=?,
                           status='scheduled', updated_at=CURRENT_TIMESTAMP
                       WHERE id=? AND employer_id=?""",
                    (
                        scheduled_at, mode, meeting_url, location, notes,
                        existing["id"], session["user_id"],
                    ),
                )
                notification_title = "Interview rescheduled"
            else:
                db.execute(
                    """INSERT INTO interviews
                       (application_id, employer_id, student_id, scheduled_at, mode,
                        meeting_url, location, notes)
                       VALUES(?,?,?,?,?,?,?,?)""",
                    (
                        application_id, session["user_id"], application["student_id"],
                        scheduled_at, mode, meeting_url, location, notes,
                    ),
                )
                notification_title = "Interview scheduled"
            db.execute(
                "INSERT INTO notifications(user_id, title, body, kind) VALUES(?,?,?,'interview')",
                (
                    application["student_id"],
                    notification_title,
                    f"{application['vacancy_title']}: {scheduled_at.replace('T', ' ')} ({mode}).",
                ),
            )
            db.commit()
        except sqlite3.DatabaseError:
            db.rollback()
            current_app.logger.exception("Interview schedule could not be saved")
            flash("The interview could not be saved. Please try again.", "error")
            return redirect(url_for("employer.schedule_interview", application_id=application_id))

        flash(notification_title + " successfully.", "success")
        return redirect(url_for("employer.schedule_interview", application_id=application_id))

    interview = db.execute(
        "SELECT * FROM interviews WHERE application_id=? AND employer_id=?",
        (application_id, session["user_id"]),
    ).fetchone()
    return render_template(
        "employer/schedule-interview.html",
        application=application,
        interview=interview,
        minimum_time=(datetime.now() + timedelta(minutes=5)).strftime("%Y-%m-%dT%H:%M"),
    )


@employer_bp.post("/applications/<int:application_id>/interview/cancel")
@role_required("employer")
def cancel_interview(application_id):
    db = get_db()
    item = db.execute(
        """SELECT i.id, i.status, i.student_id, v.title
           FROM interviews i
           JOIN applications a ON a.id=i.application_id
           JOIN vacancies v ON v.id=a.vacancy_id
           WHERE i.application_id=? AND v.employer_id=?""",
        (application_id, session["user_id"]),
    ).fetchone()
    if not item:
        return render_template("404.html"), 404
    if item["status"] != "scheduled":
        flash("This interview is not currently scheduled.", "error")
        return redirect(url_for("employer.schedule_interview", application_id=application_id))
    try:
        db.execute(
            "UPDATE interviews SET status='cancelled', updated_at=CURRENT_TIMESTAMP "
            "WHERE id=? AND employer_id=?",
            (item["id"], session["user_id"]),
        )
        db.execute(
            "INSERT INTO notifications(user_id, title, body, kind) VALUES(?,?,?,'interview')",
            (item["student_id"], "Interview cancelled", f"The interview for {item['title']} was cancelled."),
        )
        db.commit()
    except sqlite3.DatabaseError:
        db.rollback()
        current_app.logger.exception("Interview cancellation failed")
        flash("The interview could not be cancelled. Please try again.", "error")
    else:
        flash("Interview cancelled.", "success")
    return redirect(url_for("employer.schedule_interview", application_id=application_id))


@employer_bp.get("/applications/<int:application_id>/resume")
@role_required("employer")
def application_resume(application_id):
    db = get_db()
    row = db.execute(
        """SELECT sp.resume_filename
           FROM applications a
           JOIN vacancies v ON v.id=a.vacancy_id
           JOIN student_profiles sp ON sp.user_id=a.student_id
           WHERE a.id=? AND v.employer_id=?""",
        (application_id, session["user_id"]),
    ).fetchone()
    if not row or not row["resume_filename"]:
        return render_template("404.html"), 404

    response = send_from_directory(
        current_app.config["UPLOAD_FOLDER"],
        row["resume_filename"],
        as_attachment=True,
    )
    response.headers["Cache-Control"] = "private, no-store"
    response.headers["X-Content-Type-Options"] = "nosniff"
    return response


@employer_bp.post("/applications/<int:application_id>/status")
@role_required("employer")
def application_status(application_id):
    status = request.form.get("status")
    if status not in APPLICATION_STATUSES:
        return redirect(url_for("employer.applications"))

    db = get_db()
    row = db.execute(
        "SELECT a.id FROM applications a "
        "JOIN vacancies v ON v.id=a.vacancy_id "
        "WHERE a.id=? AND v.employer_id=?",
        (application_id, session["user_id"]),
    ).fetchone()

    if not row:
        return redirect(url_for("employer.applications"))

    db.execute(
        "UPDATE applications SET status=? WHERE id=?",
        (status, application_id),
    )
    db.commit()
    return redirect(url_for("employer.applications"))

@employer_bp.route("/password", methods=["GET", "POST"])
@role_required("employer")
def change_password():
    db = get_db()
    uid = session["user_id"]

    if request.method == "POST":
        current_password = request.form.get("current_password", "")
        new_password = request.form.get("new_password", "")
        confirm_password = request.form.get("confirm_password", "")

        row = db.execute(
            "SELECT password_hash FROM users WHERE id=? AND role='employer'",
            (uid,),
        ).fetchone()

        from werkzeug.security import check_password_hash, generate_password_hash

        if not row or not check_password_hash(row["password_hash"], current_password):
            flash("Current password is incorrect.", "error")
            return render_template("employer/change-password.html")

        if len(new_password) < 12:
            flash("New password must be at least 12 characters.", "error")
            return render_template("employer/change-password.html")

        if new_password != confirm_password:
            flash("New password and confirmation do not match.", "error")
            return render_template("employer/change-password.html")

        if new_password == current_password:
            flash("Choose a new password different from the current password.", "error")
            return render_template("employer/change-password.html")

        new_auth_version = secrets.token_urlsafe(32)
        db.execute(
            "UPDATE users SET password_hash=?, auth_version=? WHERE id=? AND role='employer'",
            (generate_password_hash(new_password), new_auth_version, uid),
        )
        db.commit()
        # Keep this browser signed in while invalidating all other sessions.
        session["auth_version"] = new_auth_version
        flash("Password changed successfully.", "success")
        return redirect(url_for("employer.profile"))

    return render_template("employer/change-password.html")

